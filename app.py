"""Single-server question practice application; public API owns all state changes."""
import csv
import copy
import hashlib
import io
import json
import os
import re
import random
import secrets
import sqlite3
import time
import zipfile
from contextlib import contextmanager, asynccontextmanager
from pathlib import Path

from argon2 import PasswordHasher
from argon2.exceptions import VerificationError
from fastapi import FastAPI, HTTPException, Request, Response, UploadFile, File
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook

DATABASE = os.getenv('DATABASE', 'data/practice.db')
SECURE = os.getenv('COOKIE_SECURE', '0') == '1'
ph = PasswordHasher()
@asynccontextmanager
async def lifespan(application):
    initialize()
    yield


app = FastAPI(title='亳州网络部在线刷题', lifespan=lifespan)


@contextmanager
def transaction():
    db = sqlite3.connect(DATABASE, timeout=20)
    db.row_factory = sqlite3.Row
    db.execute('PRAGMA foreign_keys=ON')
    db.execute('BEGIN IMMEDIATE')
    try:
        yield db
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def initialize():
    Path(DATABASE).parent.mkdir(parents=True, exist_ok=True)
    with transaction() as db:
        db.executescript('''
        CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY, username TEXT UNIQUE, nickname TEXT, hash TEXT, role TEXT, must_change INTEGER DEFAULT 0);
        CREATE TABLE IF NOT EXISTS logins(token TEXT PRIMARY KEY, user_id TEXT REFERENCES users(id), expires REAL);
        CREATE TABLE IF NOT EXISTS objects(kind TEXT, id TEXT, owner TEXT, body TEXT NOT NULL, PRIMARY KEY(kind,id));
        CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY, actor TEXT, action TEXT, target TEXT, at REAL);
        CREATE TABLE IF NOT EXISTS attempts(key TEXT PRIMARY KEY, count INTEGER, start REAL);
        ''')
        columns = {r['name'] for r in db.execute('PRAGMA table_info(users)')}
        for name, definition in [('name', 'TEXT'), ('created_at', 'REAL'), ('disabled', 'INTEGER NOT NULL DEFAULT 0')]:
            if name not in columns:
                db.execute(f'ALTER TABLE users ADD COLUMN {name} {definition}')
        if not db.execute("SELECT 1 FROM users WHERE role='admin'").fetchone():
            password = os.getenv('ADMIN_PASSWORD', '')
            if password:
                validate_password(password)
                name = os.getenv('ADMIN_USERNAME', 'admin').lower()
                if not re.fullmatch(r'[a-z0-9_]{3,32}', name):
                    raise ValueError('管理员用户名需为3–32位字母、数字或下划线')
                db.execute('INSERT INTO users(id,username,nickname,hash,role,created_at) VALUES(?,?,?,?,?,?)', (newid(), name, '管理员', ph.hash(password), 'admin', time.time()))


def newid():
    return secrets.token_hex(16)


def fail(code, message):
    raise HTTPException(code, message)


def validate_password(password):
    if not isinstance(password, str) or not 10 <= len(password) <= 128:
        fail(400, '密码长度需为 10–128 个字符')


def put(db, kind, item, owner=''):
    db.execute('INSERT INTO objects VALUES(?,?,?,?) ON CONFLICT(kind,id) DO UPDATE SET body=excluded.body', (kind, item['id'], owner, json.dumps(item, ensure_ascii=False)))


def get(db, kind, id):
    row = db.execute('SELECT body FROM objects WHERE kind=? AND id=?', (kind, id)).fetchone()
    if not row:
        fail(404, '记录不存在')
    return json.loads(row[0])


def items(db, kind, owner=None):
    rows = db.execute('SELECT body FROM objects WHERE kind=?' + (' AND owner=?' if owner is not None else ''), (kind, owner) if owner is not None else (kind,)).fetchall()
    return [json.loads(r[0]) for r in rows]


def user(db, request, admin=False, allow_change=False):
    token = hashlib.sha256(request.cookies.get('session', '').encode()).hexdigest()
    u = db.execute('SELECT users.* FROM users JOIN logins ON users.id=logins.user_id WHERE token=? AND expires>?', (token, time.time())).fetchone()
    if not u:
        fail(401, '请先登录')
    if u['disabled']:
        fail(403, '账号已停用，请联系管理员')
    if not allow_change and u['must_change']:
        fail(403, '请先修改临时密码')
    if admin and u['role'] != 'admin':
        fail(403, '仅管理员可操作')
    return dict(u)


def public_user(u):
    return {k: u[k] for k in ('id', 'username', 'nickname', 'role', 'must_change', 'name', 'created_at', 'disabled')}


@app.middleware('http')
async def security(request, call_next):
    if request.method not in ('GET', 'HEAD', 'OPTIONS'):
        origin = request.headers.get('origin')
        expected = os.getenv('PUBLIC_ORIGIN') or str(request.base_url).rstrip('/')
        if origin and origin != expected:
            return JSONResponse({'detail': '请求来源不允许'}, status_code=403)
        if request.headers.get('sec-fetch-site') == 'cross-site':
            return JSONResponse({'detail': '不允许跨站操作'}, status_code=403)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Content-Security-Policy'] = "default-src 'self'; style-src 'self' 'unsafe-inline'; script-src 'self'; frame-ancestors 'none'; base-uri 'self'"
    if request.url.path.startswith('/api/'):
        response.headers['Cache-Control'] = 'no-store'
    return response


def auth_limit(db, request):
    key = request.client.host if request.client else 'unknown'
    now = time.time()
    db.execute('DELETE FROM attempts WHERE start<?', (now - 900,))
    row = db.execute('SELECT * FROM attempts WHERE key=?', (key,)).fetchone()
    if row and row['count'] >= 40:
        fail(429, '尝试过于频繁，请稍后再试')
    db.execute('INSERT INTO attempts VALUES(?,1,?) ON CONFLICT(key) DO UPDATE SET count=count+1', (key, now))


@app.post('/api/register')
def register(body: dict, request: Request):
    name, password = body.get('username', ''), body.get('password')
    if not isinstance(name, str) or not re.fullmatch(r'[A-Za-z0-9_]{3,32}', name):
        fail(400, '用户名为 3–32 位字母、数字或下划线')
    validate_password(password)
    real_name = body.get('name')
    if not isinstance(real_name, str) or not 1 <= len(real_name.strip()) <= 50:
        fail(400, '姓名需为 1–50 个字符')
    with transaction() as db:
        auth_limit(db, request)
    with transaction() as db:
        try:
            db.execute('INSERT INTO users(id,username,nickname,hash,role,name,created_at) VALUES(?,?,?,?,?,?,?)', (newid(), name.lower(), real_name.strip(), ph.hash(password), 'user', real_name.strip(), time.time()))
        except sqlite3.IntegrityError:
            fail(409, '用户名已存在')
    return {'ok': True}


def issue_login(db, response, uid):
    raw = secrets.token_urlsafe(32)
    db.execute('DELETE FROM logins WHERE expires<?', (time.time(),))
    db.execute('INSERT INTO logins VALUES(?,?,?)', (hashlib.sha256(raw.encode()).hexdigest(), uid, time.time() + 604800))
    response.set_cookie('session', raw, httponly=True, samesite='strict', secure=SECURE, max_age=604800)


@app.post('/api/login')
def login(body: dict, request: Request, response: Response):
    with transaction() as db:
        auth_limit(db, request)
    with transaction() as db:
        u = db.execute('SELECT * FROM users WHERE username=?', (str(body.get('username', '')).lower(),)).fetchone()
        try:
            if not u or not ph.verify(u['hash'], str(body.get('password', ''))[:128]):
                fail(401, '用户名或密码错误')
        except VerificationError:
            fail(401, '用户名或密码错误')
        if u['disabled']:
            fail(403, '账号已停用，请联系管理员')
        if body.get('role') is not None and body['role'] != u['role']:
            fail(403, '账号角色与所选身份不符，请切换登录身份后重试')
        issue_login(db, response, u['id'])
        return public_user(u)


@app.post('/api/logout')
def logout(request: Request, response: Response):
    with transaction() as db:
        db.execute('DELETE FROM logins WHERE token=?', (hashlib.sha256(request.cookies.get('session', '').encode()).hexdigest(),))
    response.delete_cookie('session')
    return {'ok': True}


@app.get('/api/me')
def me(request: Request):
    with transaction() as db:
        return public_user(user(db, request, allow_change=True))


@app.post('/api/password')
def password(body: dict, request: Request, response: Response):
    validate_password(body.get('password'))
    with transaction() as db:
        u = user(db, request, allow_change=True)
        try:
            ph.verify(u['hash'], str(body.get('old', '')))
        except VerificationError:
            fail(400, '原密码错误')
        db.execute('UPDATE users SET hash=?,must_change=0 WHERE id=?', (ph.hash(body['password']), u['id']))
        db.execute('DELETE FROM logins WHERE user_id=?', (u['id'],))
        issue_login(db, response, u['id'])
    return {'ok': True}


@app.get('/api/users')
def users(request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        return [public_user(r) for r in db.execute('SELECT * FROM users').fetchall()]


@app.get('/api/admin/summary')
def admin_summary(request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        all_banks = items(db, 'bank')
        return {'banks': len(all_banks), 'published': sum(b['status'] == 'published' for b in all_banks),
                'drafts': sum(b['draft'] is not None or b['status'] == 'draft' for b in all_banks),
                'learners': db.execute("SELECT COUNT(*) FROM users WHERE role='user'").fetchone()[0],
                'pending_reports': sum(r['status'] == '待处理' for r in items(db, 'report'))}


@app.post('/api/users/{uid}/disable')
def disable_user(uid: str, request: Request):
    with transaction() as db:
        actor = user(db, request, admin=True)
        target = db.execute('SELECT * FROM users WHERE id=?', (uid,)).fetchone()
        if not target:
            fail(404, '用户不存在')
        if target['role'] != 'user':
            fail(400, '只能停用普通用户账号')
        db.execute('UPDATE users SET disabled=1 WHERE id=?', (uid,))
        db.execute('DELETE FROM logins WHERE user_id=?', (uid,))
        if not target['disabled']:
            db.execute('INSERT INTO audit(actor,action,target,at) VALUES(?,?,?,?)', (actor['id'], 'disable_user', uid, time.time()))
    return {'ok': True}


@app.post('/api/users/{uid}/reset')
def reset(uid: str, body: dict, request: Request):
    validate_password(body.get('password'))
    with transaction() as db:
        u = user(db, request, admin=True)
        if not db.execute('SELECT 1 FROM users WHERE id=?', (uid,)).fetchone():
            fail(404, '用户不存在')
        db.execute('UPDATE users SET hash=?,must_change=1 WHERE id=?', (ph.hash(body['password']), uid))
        db.execute('DELETE FROM logins WHERE user_id=?', (uid,))
        db.execute('INSERT INTO audit(actor,action,target,at) VALUES(?,?,?,?)', (u['id'], 'reset_password', uid, time.time()))
    return {'ok': True}


def directory_name(body):
    name = body.get('name')
    if not isinstance(name, str) or not 1 <= len(name.strip()) <= 100:
        fail(400, '目录名称需为 1–100 个字符')
    return name.strip()


def directory_reference(db, kind, value):
    if not isinstance(value, str) or not value:
        fail(400, '归属标识无效')
    return get(db, kind, value)


@app.get('/api/directories')
def directories(request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        return {'years': items(db, 'year'), 'projects': items(db, 'project')}


@app.get('/api/learning/directories')
def learning_directories(request: Request):
    with transaction() as db:
        user(db, request)
        project_ids = {b.get('project_id') for b in items(db, 'bank') if bank_available(b)}
        projects = [p for p in items(db, 'project') if p['id'] in project_ids]
        year_ids = {p['year_id'] for p in projects}
        return {'years': [y for y in items(db, 'year') if y['id'] in year_ids], 'projects': projects}


@app.post('/api/years')
def create_year(body: dict, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        year = {'id': newid(), 'name': directory_name(body)}
        put(db, 'year', year)
        return year


@app.patch('/api/years/{year_id}')
def rename_year(year_id: str, body: dict, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        year = get(db, 'year', year_id)
        year['name'] = directory_name(body)
        put(db, 'year', year)
        return year


@app.delete('/api/years/{year_id}')
def delete_year(year_id: str, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        get(db, 'year', year_id)
        if any(p['year_id'] == year_id for p in items(db, 'project')):
            fail(409, '年份下仍有项目，只能删除空目录')
        db.execute("DELETE FROM objects WHERE kind='year' AND id=?", (year_id,))
        return {'ok': True}


@app.post('/api/projects')
def create_project(body: dict, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        year = directory_reference(db, 'year', body.get('year_id'))
        project = {'id': newid(), 'name': directory_name(body), 'year_id': year['id']}
        put(db, 'project', project)
        return project


@app.patch('/api/projects/{project_id}')
def rename_project(project_id: str, body: dict, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        project = get(db, 'project', project_id)
        if 'name' in body:
            project['name'] = directory_name(body)
        if 'hidden' in body:
            if not isinstance(body['hidden'], bool):
                fail(400, '隐藏状态需为布尔值')
            project['hidden'] = body['hidden']
        if not {'name', 'hidden'} & body.keys():
            fail(400, '请提供名称或隐藏状态')
        put(db, 'project', project)
        return project


@app.delete('/api/projects/{project_id}')
def delete_project(project_id: str, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        get(db, 'project', project_id)
        if any(b.get('project_id') == project_id for b in items(db, 'bank')):
            fail(409, '项目下仍有题库，只能删除空目录')
        db.execute("DELETE FROM objects WHERE kind='project' AND id=?", (project_id,))
        return {'ok': True}


@app.get('/api/banks')
def banks(request: Request):
    with transaction() as db:
        u = user(db, request)
        return [bank_view(b, u) for b in items(db, 'bank') if u['role'] == 'admin' or bank_available(b)]


@app.post('/api/banks')
def create_bank(body: dict, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        title = str(body.get('title', '')).strip()
        if not title or len(title) > 100:
            fail(400, '题库名称需为 1–100 个字符')
        b = {'id': newid(), 'title': title, 'status': 'draft', 'revision': 0, 'questions': [], 'draft': None}
        if 'project_id' in body:
            b['project_id'] = directory_reference(db, 'project', body['project_id'])['id']
        put(db, 'bank', b)
        return b


def bank_available(bank):
    # Missing ownership is the legacy pending area; never infer it from titles.
    return bank['status'] == 'published' and bool(bank.get('project_id'))


def bank_view(b, u):
    if u['role'] == 'admin':
        return b
    active = [q for q in b['questions'] if not q.get('archived')]
    return {k: b[k] for k in ('id', 'title', 'status', 'revision')} | {'project_id': b.get('project_id'), 'count': len(active), 'categories': sorted({q['category'] for q in active if q['category']})}


HEADERS = ['编号', '题型', '题干', 'A', 'B', 'C', 'D', 'E', '答案', '解析', '分类', '标签']
EXAMPLES = [['Q1', '单选', '一加一等于？', '一', '二', '三', '', 'B', '一加一等于二。', '基础', '数学'], ['Q2', '多选', '哪些是偶数？', '二', '三', '四', '', 'A|C', '二和四是偶数。', '基础', '数学'], ['Q3', '判断', '地球是平的。', '', '', '', '', '错误', '地球不是平的。', '地理', '常识']]

LEGACY_HEADERS = [h for h in HEADERS if h != 'E']
EXAMPLES = [row[:7] + [''] + row[7:] for row in EXAMPLES]

@app.get('/api/template.{format}')
def template(format: str, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
    if format == 'csv':
        output = io.StringIO()
        csv.writer(output).writerows([HEADERS] + EXAMPLES)
        return Response(output.getvalue().encode('utf-8-sig'), media_type='text/csv', headers={'Content-Disposition': 'attachment; filename="template.csv"'})
    if format == 'xlsx':
        w = Workbook()
        for row in [HEADERS] + EXAMPLES:
            w.active.append(row)
        w.active.freeze_panes = 'A2'
        binary_output = io.BytesIO()
        w.save(binary_output)
        return Response(binary_output.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="template.xlsx"'})
    fail(400, '仅支持 CSV/XLSX')


def parse_file(raw, filename, bid):
    try:
        if filename.lower().endswith('.csv'):
            rows = list(csv.reader(io.StringIO(raw.decode('utf-8-sig')), strict=True))
        elif filename.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50_000_000:
                    fail(400, '解压后文件过大')
            w = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            rows = []
            for row in w.active.iter_rows():
                if any(c.data_type == 'f' for c in row):
                    fail(400, '模板不允许公式，请转换为文本')
                rows.append([str(c.value) if c.value is not None else '' for c in row])
                if len(rows) > 10001:
                    fail(400, '最多导入 10000 题')
            w.close()
        else:
            fail(400, '仅支持 UTF-8 CSV 或 XLSX')
    except HTTPException:
        raise
    except Exception:
        fail(400, '无法读取文件，请使用下载的模板和 UTF-8 编码')
    if not rows or rows[0] not in (HEADERS, LEGACY_HEADERS):
        fail(400, '表头必须与模板完全一致')
    if len(rows) > 10001:
        fail(400, '最多导入 10000 题')
    errors, questions, codes = [], [], set()
    for line, row in enumerate(rows[1:], 2):
        if not any(row):
            continue
        if len(row) != len(rows[0]):
            errors.append({'row': line, 'field': '列数', 'message': '列数与模板不一致'})
            continue
        r = dict(zip(rows[0], [str(x).strip() for x in row]))
        r.setdefault('E', '')
        problems = []
        if not r['编号'] or len(r['编号']) > 64 or r['编号'] in codes:
            problems.append(('编号', '编号必填、最多64字符且文件内不能重复'))
        codes.add(r['编号'])
        if not r['题干'] or len(r['题干']) > 10000:
            problems.append(('题干', '题干必填且最多10000字符'))
        if r['题型'] not in ('单选', '多选', '判断'):
            problems.append(('题型', '仅支持单选、多选、判断'))
        options = {key: r[key] for key in 'ABCDE' if r[key]}
        answer = r['答案'].split('|')
        if r['题型'] == '判断':
            options = {'A': '正确', 'B': '错误'}
            answer = [{'正确': 'A', '错误': 'B'}.get(r['答案'], '?')]
        if len(options) < 2 or not set(answer).issubset(options) or len(answer) != len(set(answer)) or (r['题型'] != '多选' and len(answer) != 1):
            problems.append(('答案', '答案需引用有效选项；单选一个答案，多选用 | 分隔，判断填正确或错误'))
        if any(len(v) > 10000 for v in r.values()):
            problems.append(('内容', '单个字段最多10000字符'))
        errors.extend({'row': line, 'field': field, 'message': message} for field, message in problems)
        questions.append({'id': bid + ':' + r['编号'], 'code': r['编号'], 'type': r['题型'], 'stem': r['题干'], 'options': options, 'answer': sorted(answer), 'explanation': r['解析'], 'category': r['分类'], 'tags': [t for t in r['标签'].split('|') if t], 'archived': False})
    if not questions and not errors:
        errors.append({'row': 2, 'field': '题目', 'message': '文件没有题目'})
    return questions, errors


@app.post('/api/banks/{bid}/preview')
async def preview(bid: str, request: Request, file: UploadFile = File(...)):
    raw = await file.read(10_000_001)
    if len(raw) > 10_000_000:
        fail(400, '文件最多10 MB')
    with transaction() as db:
        u = user(db, request, admin=True)
        b = get(db, 'bank', bid)
        questions, errors = parse_file(raw, file.filename or '', bid)
        base = b['draft'] if b['draft'] is not None else b['questions']
        known = {q['code'] for q in base}
        p = {'id': newid(), 'bank_id': bid, 'revision': b['revision'], 'questions': questions, 'errors': errors, 'added': sum(q['code'] not in known for q in questions), 'updated': sum(q['code'] in known for q in questions), 'committed': False}
        if not errors:
            put(db, 'import', p, u['id'])
        return p


@app.post('/api/imports/{iid}/commit')
def commit_import(iid: str, request: Request, body: dict | None = None):
    publish_now = (body or {}).get("publish", False)
    if not isinstance(publish_now, bool):
        fail(400, "publish 必须为布尔值")
    with transaction() as db:
        user(db, request, admin=True)
        p = get(db, 'import', iid)
        if p['committed']:
            if p.get('published', False) != publish_now:
                fail(409, '该预览已确认，不能更改生效方式，请重新预览')
            return {'ok': True}
        b = get(db, 'bank', p['bank_id'])
        check_revision(b, p)
        base = {q['code']: q for q in (b['draft'] if b['draft'] is not None else b['questions'])}
        for q in p['questions']:
            base[q['code']] = q
        b['draft'] = list(base.values())
        if publish_now:
            publish_content(db, b)
        b['revision'] += 1
        put(db, 'bank', b)
        p['committed'] = True
        p['published'] = publish_now
        put(db, 'import', p)
        return {'ok': True}


def check_revision(item, body):
    if body.get('revision') != item['revision']:
        fail(409, '数据已在另一设备修改，请刷新后重试；本地选择已保留')


def publish_content(db, b):
    old = {q['id']: q for q in b['questions']}
    if b['draft'] is not None:
        for q in b['draft']:
            previous = old.get(q['id'])
            if previous and previous['answer'] != q['answer']:
                for learning in items(db, 'learning'):
                    if q['id'] in learning['mistakes']:
                        learning['mistakes'][q['id']]['mastered'] = False
                        put(db, 'learning', learning, learning['id'])
            q['version'] = (previous or {}).get('version', 0) + 1
        b['questions'] = b['draft']
        b['draft'] = None
    if not b['questions']:
        fail(400, '请先导入题目')
    b['status'] = 'published'


@app.post('/api/banks/{bid}/publish')
def publish(bid: str, body: dict, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        b = get(db, 'bank', bid)
        check_revision(b, body)
        publish_content(db, b)
        b['revision'] += 1
        put(db, 'bank', b)
        return b


@app.patch('/api/banks/{bid}')
def edit_bank(bid: str, body: dict, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        b = get(db, 'bank', bid)
        check_revision(b, body)
        if 'title' in body:
            title = str(body['title']).strip()
            if not title or len(title) > 100:
                fail(400, '名称需为1–100字符')
            b['title'] = title
        if 'project_id' in body:
            b['project_id'] = directory_reference(db, 'project', body['project_id'])['id']
        if 'status' in body:
            if body['status'] not in ('offline', 'archived'):
                fail(400, '此入口只支持下架或归档')
            b['status'] = body['status']
        b['revision'] += 1
        put(db, 'bank', b)
        return b


@app.post('/api/banks/{bid}/archive')
def archive(bid: str, body: dict, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        b = get(db, 'bank', bid)
        check_revision(b, body)
        draft = copy.deepcopy(b['draft'] if b['draft'] is not None else b['questions'])
        q = next((q for q in draft if q['id'] == body.get('question_id')), None)
        if not q:
            fail(404, '题目不存在')
        assert q is not None
        q['archived'] = True
        b['draft'] = draft
        b['revision'] += 1
        put(db, 'bank', b)
        return b


@app.post('/api/banks/{bid}/archive-batch')
def archive_batch(bid: str, body: dict, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        bank = get(db, 'bank', bid)
        check_revision(bank, body)
        draft = copy.deepcopy(bank['draft'] if bank['draft'] is not None else bank['questions'])
        if body.get('all') is True:
            if body.get('question_ids'):
                fail(400, '全部删除不能同时指定题目')
            ids = {q['id'] for q in draft if not q.get('archived')}
        else:
            selected = body.get('question_ids')
            if not isinstance(selected, list) or not selected or len(selected) > 10000 or any(not isinstance(id, str) for id in selected):
                fail(400, '请选择要删除的题目')
            assert isinstance(selected, list)
            ids = set(selected)
            if not ids.issubset({q['id'] for q in draft}):
                fail(400, '选择包含不存在的题目，请刷新后重试')
        count = 0
        for q in draft:
            if q['id'] in ids and not q.get('archived'):
                q['archived'] = True
                count += 1
        if not count:
            fail(400, '没有可删除的题目')
        bank['draft'] = draft
        bank['revision'] += 1
        put(db, 'bank', bank)
        return {'archived': count, 'revision': bank['revision']}


def learning_state(db, uid):
    row = db.execute("SELECT body FROM objects WHERE kind='learning' AND id=?", (uid,)).fetchone()
    return json.loads(row[0]) if row else {'id': uid, 'mistakes': {}, 'favorites': [], 'cursors': {}}


def owned_session(db, sid, uid):
    s = get(db, 'session', sid)
    if s['user_id'] != uid:
        fail(404, '练习不存在')
    expire_session(db, s)
    return s


def expire_session(db, s):
    deadline = s.get('deadline')
    if not s['ended'] and deadline is not None and time.time() >= deadline:
        s['ended'] = True
        s['ended_at'] = deadline
        s['revision'] += 1
        put(db, 'session', s, s['user_id'])


def timed_session_closed(s):
    if s.get('deadline') is not None and s['ended']:
        # Return normally so transaction() commits automatic expiration even on rejection.
        return JSONResponse({'detail': '限时练习已结束，未提交题目不计分', 'session': session_view(s)}, status_code=409)
    return None


def session_view(s):
    result = dict(s)
    if s.get('deadline') is not None:
        result['server_now'] = time.time()
    result['questions'] = [{k: v for k, v in q.items() if k not in ('answer', 'explanation') or str(i) in s['answers']} for i, q in enumerate(s['questions'])]
    return result


def session_project(db, s):
    if 'project' in s:
        return s['project']
    bank = get(db, 'bank', s['bank_id'])
    if not bank.get('project_id'):
        return None
    project = get(db, 'project', bank['project_id'])
    year = get(db, 'year', project['year_id'])
    return {'id': project['id'], 'name': project['name'], 'year_id': year['id'], 'year_name': year['name']}


@app.get('/api/sessions')
def sessions(request: Request):
    with transaction() as db:
        u = user(db, request)
        rows = sorted(items(db, 'session', u['id']), key=lambda s: s['created'], reverse=True)
        for s in rows:
            expire_session(db, s)
        return [{k: s[k] for k in ('id', 'title', 'mode', 'ended', 'created', 'index')} | {'project': session_project(db, s), 'answered': len(s['answers']), 'count': len(s['questions']), 'correct': sum(a['correct'] for a in s['answers'].values())} for s in rows]


@app.post('/api/sessions')
def create_session(body: dict, request: Request):
    with transaction() as db:
        u = user(db, request)
        b = get(db, 'bank', body.get('bank_id', ''))
        if not bank_available(b):
            fail(400, '题库待整理、尚未发布、已下架或已归档，不能开始新练习')
        mode = body.get('mode', '顺序练习')
        if mode not in ('顺序练习', '随机练习', '错题练习', '收藏练习') or str(body.get('amount', '20')) not in ('10', '20', '50', '全部'):
            fail(400, '无效练习配置')
        review_mode = body.get('review_mode', 'immediate')
        if review_mode not in ('immediate', 'deferred'):
            fail(400, '无效核对方式')
        duration = body.get('duration_minutes')
        if duration is not None and (type(duration) is not int or not 1 <= duration <= 1440):
            fail(400, '练习时长需为 1–1440 的整数分钟；不限时请留空')
        state = learning_state(db, u['id'])
        category = str(body.get('category', ''))
        scope = b['id'] + '/' + category
        qs = [q for q in b['questions'] if not q.get('archived') and (not category or q['category'] == category)]
        if mode == '错题练习':
            qs = [q for q in qs if q['id'] in state['mistakes'] and not state['mistakes'][q['id']]['mastered']]
        elif mode == '收藏练习':
            qs = [q for q in qs if q['id'] in state['favorites']]
        elif mode == '随机练习':
            random.SystemRandom().shuffle(qs)
        elif not body.get('restart'):
            cursor = state['cursors'].get(scope)
            if cursor:
                location = next((i for i, q in enumerate(qs) if q['id'] == cursor), -1)
                if location >= 0:
                    qs = qs[location + 1:]
        if not qs:
            fail(400, '此范围没有可练习题目；顺序已完成时可选择从头开始')
        amount = str(body.get('amount', '20'))
        if amount != '全部':
            qs = qs[:int(amount)]
        s = {'id': newid(), 'user_id': u['id'], 'bank_id': b['id'], 'title': b['title'], 'mode': mode, 'review_mode': review_mode, 'scope': scope, 'questions': qs, 'drafts': {}, 'answers': {}, 'index': 0, 'revision': 0, 'ended': False, 'created': time.time()}
        if duration is not None:
            s['deadline'] = s['created'] + duration * 60
        s['project'] = session_project(db, s)
        put(db, 'session', s, u['id'])
        return session_view(s)


@app.get('/api/sessions/{sid}')
def session(sid: str, request: Request):
    with transaction() as db:
        u = user(db, request)
        return session_view(owned_session(db, sid, u['id']))


def selection(s, body):
    index = body.get('index')
    if not isinstance(index, int) or isinstance(index, bool) or not 0 <= index < len(s['questions']):
        fail(400, '题目位置无效')
    selected = body.get('selected', [])
    if not isinstance(selected, list) or any(not isinstance(x, str) for x in selected) or len(selected) != len(set(selected)) or not set(selected).issubset(s['questions'][index]['options']):
        fail(400, '选项无效')
    if s['questions'][index]['type'] != '多选' and len(selected) > 1:
        fail(400, '单选和判断只能选择一项')
    return index, sorted(selected)


@app.patch('/api/sessions/{sid}')
def save_draft(sid: str, body: dict, request: Request):
    with transaction() as db:
        u = user(db, request)
        s = owned_session(db, sid, u['id'])
        closed = timed_session_closed(s)
        if closed is not None:
            return closed
        check_revision(s, body)
        if s['ended']:
            fail(400, '练习已结束')
        i, selected = selection(s, body)
        if str(i) not in s['answers']:
            s['drafts'][str(i)] = selected
        s['index'] = i
        s['revision'] += 1
        put(db, 'session', s, u['id'])
        return session_view(s)


def record_answer(db, s, u, i, selected, state=None):
    q = s['questions'][i]
    correct = selected == sorted(q['answer'])
    s['answers'][str(i)] = {'selected': selected, 'correct': correct, 'at': time.time()}
    s['drafts'].pop(str(i), None)
    s['index'] = i
    s['revision'] += 1
    shared_state = state is not None
    if state is None:
        state = learning_state(db, u['id'])
    if not correct:
        state['mistakes'][q['id']] = {'mastered': False, 'title': q['stem'], 'bank_id': s['bank_id']}
    if s['mode'] == '顺序练习' and not shared_state:
        # Advance only through the contiguous answered prefix; skipped questions remain reachable.
        for n, question in enumerate(s['questions']):
            if str(n) not in s['answers']:
                break
            state['cursors'][s['scope']] = question['id']
    if not shared_state:
        put(db, 'learning', state, u['id'])


@app.post('/api/sessions/{sid}/answer')
def submit_answer(sid: str, body: dict, request: Request):
    with transaction() as db:
        u = user(db, request)
        s = owned_session(db, sid, u['id'])
        closed = timed_session_closed(s)
        if closed is not None:
            return closed
        i, selected = selection(s, body)
        if str(i) in s['answers']:
            if s['answers'][str(i)]['selected'] == selected:
                return session_view(s)
            fail(409, '答案已提交，不能修改')
        check_revision(s, body)
        if s.get('review_mode') == 'deferred':
            fail(400, '本场练习需全部选择后统一核对')
        if s['ended'] or not selected:
            fail(400, '练习已结束或尚未选择答案')
        record_answer(db, s, u, i, selected)
        put(db, 'session', s, u['id'])
        return session_view(s)


@app.post('/api/sessions/{sid}/review')
def review_answers(sid: str, body: dict, request: Request):
    with transaction() as db:
        u = user(db, request)
        s = owned_session(db, sid, u['id'])
        closed = timed_session_closed(s)
        if closed is not None:
            return closed
        if s.get('reviewed'):
            return session_view(s)
        check_revision(s, body)
        if s['ended'] or s.get('review_mode') != 'deferred':
            fail(400, '当前练习不能统一核对')
        selections = [selection(s, {'index': i, 'selected': s['drafts'].get(str(i), [])})[1] for i in range(len(s['questions']))]
        if any(not selected for selected in selections):
            fail(400, '请先选择全部题目的答案')
        state = learning_state(db, u['id'])
        for i, selected in enumerate(selections):
            record_answer(db, s, u, i, selected, state)
        if s['mode'] == '顺序练习':
            state['cursors'][s['scope']] = s['questions'][-1]['id']
        put(db, 'learning', state, u['id'])
        s['ended'] = True
        s['reviewed'] = True
        s['index'] = 0
        put(db, 'session', s, u['id'])
        return session_view(s)


@app.post('/api/sessions/{sid}/finish')
def finish(sid: str, body: dict, request: Request):
    with transaction() as db:
        u = user(db, request)
        s = owned_session(db, sid, u['id'])
        if not s['ended']:
            check_revision(s, body)
            s['ended'] = True
            s['revision'] += 1
            put(db, 'session', s, u['id'])
        return session_view(s)


@app.get('/api/learning')
def learning(request: Request):
    with transaction() as db:
        u = user(db, request)
        state = learning_state(db, u['id'])
        submitted = [(q, s['answers'][str(i)]) for s in items(db, 'session', u['id']) for i, q in enumerate(s['questions']) if str(i) in s['answers']]
        active = {q['id'] for b in items(db, 'bank') if bank_available(b) for q in b['questions'] if not q.get('archived')}
        covered = {q['id'] for q, _ in submitted} & active
        records: dict[str, dict] = {}
        for s in sorted(items(db, 'session', u['id']), key=lambda s: s['created']):
            for i, q in enumerate(session_view(s)['questions']):
                if q['id'] in state['mistakes'] or q['id'] in state['favorites']:
                    previous = records.get(q['id'])
                    if previous is None or str(i) in s['answers'] or 'answer' not in previous['question']:
                        records[q['id']] = {'question': q, 'session_id': s['id'], 'index': i, 'title': s['title'], 'project': session_project(db, s)}
        return state | {'records': list(records.values()), 'total': len(submitted), 'correct': sum(a['correct'] for _, a in submitted), 'coverage': len(covered), 'available': len(active)}


@app.post('/api/favorite')
def favorite(body: dict, request: Request):
    with transaction() as db:
        u = user(db, request)
        qid = body.get('question_id')
        available = any(q['id'] == qid for b in items(db, 'bank') if bank_available(b) for q in b['questions']) or any(q['id'] == qid for s in items(db, 'session', u['id']) for q in s['questions'])
        if not available:
            fail(404, '题目不可访问')
        state = learning_state(db, u['id'])
        state['favorites'] = [q for q in state['favorites'] if q != qid]
        if body.get('value'):
            state['favorites'].append(qid)
        put(db, 'learning', state, u['id'])
        return {'ok': True}


@app.post('/api/mastered')
def mastered(body: dict, request: Request):
    with transaction() as db:
        u = user(db, request)
        state = learning_state(db, u['id'])
        if body.get('question_id') not in state['mistakes']:
            fail(404, '错题不存在')
        state['mistakes'][body['question_id']]['mastered'] = bool(body.get('value'))
        put(db, 'learning', state, u['id'])
        return {'ok': True}


@app.post('/api/reports')
def report(body: dict, request: Request):
    with transaction() as db:
        u = user(db, request)
        s = owned_session(db, body.get('session_id', ''), u['id'])
        i = body.get('index')
        if not isinstance(i, int) or not 0 <= i < len(s['questions']) or str(i) not in s['answers']:
            fail(400, '请作答后反馈')
        if body.get('type') not in ('答案有误', '解析有误', '题干有误', '选项有误'):
            fail(400, '反馈类型无效')
        r = {'id': newid(), 'user_id': u['id'], 'question': s['questions'][i], 'type': body['type'], 'text': str(body.get('text', ''))[:2000], 'status': '待处理', 'at': time.time()}
        put(db, 'report', r, u['id'])
        return {'ok': True}


@app.get('/api/reports')
def reports(request: Request):
    with transaction() as db:
        u = user(db, request)
        return items(db, 'report', None if u['role'] == 'admin' else u['id'])


@app.patch('/api/reports/{rid}')
def resolve_report(rid: str, body: dict, request: Request):
    with transaction() as db:
        user(db, request, admin=True)
        r = get(db, 'report', rid)
        if body.get('status') not in ('待处理', '已处理', '不采纳'):
            fail(400, '处理状态无效')
        r['status'] = body['status']
        put(db, 'report', r)
        return r


@app.get('/')
@app.get('/admin')
@app.get('/learn')
def index():
    built = Path(__file__).parent / 'web' / 'build' / 'index.html'
    if not built.is_file():
        fail(503, '前端尚未构建，请在 frontend 目录执行 npm ci 和 npm run build')
    return FileResponse(built)


app.mount('/assets', StaticFiles(directory=Path(__file__).parent / 'web' / 'build' / 'assets', check_dir=False), name='assets')
